"""Blackhole Tensix NoC subsystem figure — paper-ready pptx.
Every shape is a real PPT primitive so it stays editable.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---------- palette (monochromatic + one accent) ----------
INK       = RGBColor(0x1F, 0x29, 0x37)
STROKE    = RGBColor(0x6B, 0x72, 0x80)
FAINT     = RGBColor(0xD1, 0xD5, 0xDB)
BG_LIGHT  = RGBColor(0xF3, 0xF4, 0xF6)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT    = RGBColor(0xC0, 0x39, 0x2B)      # crimson
ACCENT_BG = RGBColor(0xFD, 0xEA, 0xE7)
COMPUTE   = RGBColor(0xEF, 0xF4, 0xFB)      # pale blue
MEMORY    = RGBColor(0xF3, 0xEF, 0xFA)      # pale lilac
NIU_TINT  = RGBColor(0xFA, 0xFA, 0xF9)      # very light gray

# ---------- slide ----------
prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank


# ---------- helpers ----------
def box(x, y, w, h, *, fill=WHITE, stroke=STROKE, sw=0.75, rounded=False, adj=0.06):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if stroke is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = stroke
        s.line.width = Pt(sw)
    if rounded:
        try:
            s.adjustments[0] = adj
        except Exception:
            pass
    s.shadow.inherit = False
    return s


def _apply_run(r, *, size, bold, color, name='Calibri'):
    r.font.name = name
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color


def txt(shape, text, *, size=10, bold=False, color=INK,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        pad=0.05, name='Calibri'):
    tf = shape.text_frame
    tf.margin_left = Inches(pad); tf.margin_right = Inches(pad)
    tf.margin_top = Inches(pad*0.6); tf.margin_bottom = Inches(pad*0.6)
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.text = text
    _apply_run(p.runs[0], size=size, bold=bold, color=color, name=name)
    return tf


def mtxt(shape, lines, *, sizes, bolds=None, colors=None,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, pad=0.05):
    n = len(lines)
    bolds  = bolds  or [False]*n
    colors = colors or [INK]*n
    tf = shape.text_frame
    tf.margin_left = Inches(pad); tf.margin_right = Inches(pad)
    tf.margin_top = Inches(pad*0.6); tf.margin_bottom = Inches(pad*0.6)
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        # ensure a run exists even for empty lines (used as vertical spacer)
        p.text = ln if ln else ' '
        _apply_run(p.runs[0], size=sizes[i], bold=bolds[i], color=colors[i])
    return tf


def connect(x1, y1, x2, y2, *, color=STROKE, width=0.75, arrow=False):
    from pptx.oxml.ns import qn
    from lxml import etree
    c = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
        Inches(x2 - x1), Inches(y2 - y1))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if arrow:
        ln = c.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn('a:tailEnd'))
        tail.set('type', 'triangle'); tail.set('w', 'med'); tail.set('len', 'med')
    return c


# ==========================================================
# TITLE
# ==========================================================
t = box(0.35, 0.15, 12.6, 0.55, fill=WHITE, stroke=None)
txt(t, 'Blackhole Tensix Tile — NoC subsystem',
    size=22, bold=True, align=PP_ALIGN.LEFT)


# ==========================================================
# OUTER TILE
# ==========================================================
TILE_X, TILE_Y, TILE_W, TILE_H = 0.35, 0.85, 12.63, 5.55
outer = box(TILE_X, TILE_Y, TILE_W, TILE_H,
            fill=WHITE, stroke=INK, sw=1.25, rounded=True, adj=0.015)
hdr = box(TILE_X, TILE_Y, 2.4, 0.32, fill=INK, stroke=INK, rounded=True, adj=0.04)
txt(hdr, 'Tensix Tile', size=10, bold=True, color=WHITE, align=PP_ALIGN.LEFT, pad=0.12)


# ==========================================================
# COMPUTE (5 baby RISCs) — left column
# ==========================================================
CX, CY, CW = 0.6, 1.35, 1.5
grp_lbl = box(CX, CY, CW, 0.28, fill=BG_LIGHT, stroke=FAINT, sw=0.5)
txt(grp_lbl, 'Compute (5 RISC-V)', size=9, bold=True, color=STROKE)

risc_h = 0.42; risc_gap = 0.07
risc_names = ['BRISC  (drives NoC 0)',
              'NCRISC (drives NoC 1)',
              'TRISC 0', 'TRISC 1', 'TRISC 2']
for i, n in enumerate(risc_names):
    r = box(CX, CY + 0.35 + i*(risc_h+risc_gap), CW, risc_h,
            fill=COMPUTE, stroke=STROKE, sw=0.6, rounded=True, adj=0.12)
    txt(r, n, size=9, bold=False)


# ==========================================================
# L1 SRAM — middle-left column
# ==========================================================
LX, LY, LW = 2.25, 1.35, 1.35
grp_lbl2 = box(LX, LY, LW, 0.28, fill=BG_LIGHT, stroke=FAINT, sw=0.5)
txt(grp_lbl2, 'Local memory', size=9, bold=True, color=STROKE)

l1 = box(LX, LY + 0.35, LW, 4.06, fill=MEMORY, stroke=STROKE, sw=0.6,
         rounded=True, adj=0.06)
mtxt(l1,
     ['L1 SRAM', '1.5 MB', '', 'source / target', 'of every',
      'NoC packet'],
     sizes=[14, 11, 8, 9, 9, 9],
     bolds=[True, False, False, False, False, False],
     colors=[INK, STROKE, INK, STROKE, STROKE, STROKE])


# ==========================================================
# NIU DETAIL — right side, ×2 badge
# ==========================================================
NX, NY, NW, NH = 3.85, 1.05, 8.95, 3.85
niu_hdr = box(NX, NY, NW, 0.34, fill=INK, stroke=INK, rounded=True, adj=0.03)
txt(niu_hdr,
    'NIU  (Network Interface Unit)   —   one per NoC   ×2 per tile',
    size=11, bold=True, color=WHITE, align=PP_ALIGN.LEFT, pad=0.12)
niu_body = box(NX, NY + 0.34, NW, NH - 0.34,
               fill=NIU_TINT, stroke=STROKE, sw=0.75, rounded=True, adj=0.02)

# --- CMD BUF ×4 block (top-left of NIU) ---
CB_X, CB_Y, CB_W, CB_H = NX + 0.22, NY + 0.55, 1.7, 0.95
cb = box(CB_X, CB_Y, CB_W, CB_H, fill=BG_LIGHT, stroke=STROKE, sw=0.6,
         rounded=True, adj=0.08)
mtxt(cb,
     ['CMD BUF × 4',
      '0x800 stride',
      'FIFO per buf'],
     sizes=[11, 8, 8],
     bolds=[True, False, False],
     colors=[INK, STROKE, STROKE])

# --- Register bank container ---
RX = NX + 2.05; RY = NY + 0.5
RW = NW - 2.25; RH = NH - 0.75
reg_wrap = box(RX, RY, RW, RH, fill=WHITE, stroke=FAINT, sw=0.5,
               rounded=True, adj=0.02)
reg_hdr = box(RX, RY, RW, 0.26, fill=BG_LIGHT, stroke=FAINT, sw=0.5,
              rounded=True, adj=0.06)
txt(reg_hdr,
    'Command / Status registers    (base = NOC_REGS_START_ADDR + offset)',
    size=9, bold=True, color=STROKE, align=PP_ALIGN.LEFT, pad=0.15)

# --- Register cells layout ---
# Row 1: Address (6)  |  Control (3)
# Row 2: Atomic (2)   |  ★ BRCST_EXCLUDE (highlight)
# Row 3: CMD-word bits (banner across full width)

def reg_cell(x, y, w, h, name, off, *, accent=False):
    if accent:
        f, s, sw = ACCENT_BG, ACCENT, 1.6
        color = ACCENT
    else:
        f, s, sw = WHITE, STROKE, 0.55
        color = INK
    r = box(x, y, w, h, fill=f, stroke=s, sw=sw, rounded=True, adj=0.14)
    mtxt(r, [name, off],
         sizes=[9, 7],
         bolds=[True, False],
         colors=[color, STROKE if not accent else ACCENT])
    return r


def group_wrap(x, y, w, h, title):
    # subtle wrapper with a mini header on the left
    box(x, y, w, h, fill=WHITE, stroke=FAINT, sw=0.4,
        rounded=True, adj=0.05)
    lbl = box(x + 0.05, y + 0.03, min(1.2, w - 0.1), 0.2,
              fill=BG_LIGHT, stroke=FAINT, sw=0.4, rounded=True, adj=0.15)
    txt(lbl, title, size=8, bold=True, color=STROKE, align=PP_ALIGN.LEFT, pad=0.06)


cw, ch, gap = 0.82, 0.5, 0.08

# --- Row 1 ---
row1_y = RY + 0.36
# Address group (6 cells)
addr_x = RX + 0.15
addr_w = 6*cw + 5*gap + 0.16
group_wrap(addr_x, row1_y, addr_w, ch + 0.36, 'Address')
for i, (n, o) in enumerate([('TARG_LO','0x00'),('TARG_MID','0x04'),('TARG_HI','0x08'),
                            ('RET_LO','0x0C'),('RET_MID','0x10'),('RET_HI','0x14')]):
    reg_cell(addr_x + 0.08 + i*(cw+gap), row1_y + 0.28, cw, ch, n, o)

# Row 1 continued: doesn't fit alongside — put Control on row2 first cluster instead.

# --- Row 2 ---
row2_y = row1_y + ch + 0.52
# Control (3 cells)
ctrl_x = RX + 0.15
ctrl_w = 3*cw + 2*gap + 0.16
group_wrap(ctrl_x, row2_y, ctrl_w, ch + 0.36, 'Control')
for i, (n, o) in enumerate([('NOC_CTRL','0x1C'), ('CMD_CTRL','0x40'), ('PACKET_TAG','0x18')]):
    reg_cell(ctrl_x + 0.08 + i*(cw+gap), row2_y + 0.28, cw, ch, n, o)

# Atomic (2 cells) right of Control
at_x = ctrl_x + ctrl_w + 0.12
at_w = 2*cw + 1*gap + 0.16
group_wrap(at_x, row2_y, at_w, ch + 0.36, 'Atomic')
for i, (n, o) in enumerate([('AT_LEN_BE','0x20'), ('AT_DATA','0x28')]):
    reg_cell(at_x + 0.08 + i*(cw+gap), row2_y + 0.28, cw, ch, n, o)

# ★ Multicast highlight group (BRCST_EXCLUDE only)
be_x = at_x + at_w + 0.15
be_w = 1.55
be_wrap = box(be_x, row2_y - 0.04, be_w, ch + 0.44, fill=ACCENT_BG, stroke=ACCENT, sw=1.5,
              rounded=True, adj=0.10)
be_lbl = box(be_x + 0.05, row2_y - 0.01, be_w - 0.1, 0.22,
             fill=ACCENT, stroke=ACCENT, rounded=True, adj=0.15)
txt(be_lbl, '★ Multicast (new in Blackhole)', size=8, bold=True,
    color=WHITE, align=PP_ALIGN.LEFT, pad=0.06)
be_cell = box(be_x + 0.08, row2_y + 0.28, be_w - 0.16, ch, fill=WHITE, stroke=ACCENT,
              sw=1.4, rounded=True, adj=0.12)
mtxt(be_cell,
     ['BRCST_EXCLUDE',
      '0x2C  ·  24-bit mask'],
     sizes=[10, 7], bolds=[True, False],
     colors=[ACCENT, ACCENT])

# --- Row 3: CMD-word bit strip ---
row3_y = row2_y + ch + 0.55
strip = box(RX + 0.15, row3_y, RW - 0.3, 0.48, fill=BG_LIGHT, stroke=FAINT, sw=0.5,
            rounded=True, adj=0.06)
mtxt(strip,
     ['NOC_CTRL command bits'],
     sizes=[8], bolds=[True], colors=[STROKE], anchor=MSO_ANCHOR.TOP)
# add sub line manually
p2 = strip.text_frame.add_paragraph()
p2.alignment = PP_ALIGN.CENTER
p2.text = ('BRCST_PACKET (b5)     BRCST_XY (b16)     BRCST_SRC_INCLUDE (b17)     '
           '★ combined with BRCST_EXCLUDE for per-mcast destination drop')
_apply_run(p2.runs[0], size=9, bold=False, color=INK)


# ==========================================================
# ROUTERS  (bottom row inside tile)
# ==========================================================
RTR_Y = 5.15
r0 = box(0.6, RTR_Y, 6.05, 0.6, fill=INK, stroke=INK, rounded=True, adj=0.18)
txt(r0, 'NoC 0 router  ·  X-Y routing',
    size=11, bold=True, color=WHITE)
r1 = box(6.75, RTR_Y, 6.05, 0.6, fill=INK, stroke=INK, rounded=True, adj=0.18)
txt(r1, 'NoC 1 router  ·  Y-X routing (disjoint from NoC 0)',
    size=11, bold=True, color=WHITE)


# ==========================================================
# MESH LEGEND (below tile)
# ==========================================================
mesh_y = 6.55
mesh = box(0.35, mesh_y, 12.63, 0.4, fill=WHITE, stroke=None)
txt(mesh,
    '↑ N     ↓ S     ← W     → E     2D mesh links to neighboring Tensix / DRAM / PCIe / ARC tiles',
    size=10, bold=False, color=STROKE)

# caption strip
cap_y = 6.95
cap = box(0.35, cap_y, 12.63, 0.5, fill=WHITE, stroke=None)
mtxt(cap,
     ['Fig. — BRCST_EXCLUDE is a Blackhole-generation register (absent in Wormhole) that lets a single multicast '
      'issue drop a bitmask of destinations inside the target rectangle. HW & firmware expose it; the kernel-facing '
      'dataflow API does not yet wire it through.'],
     sizes=[9], bolds=[False], colors=[STROKE],
     align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, pad=0.05)


# ==========================================================
# CONNECTORS
# ==========================================================
# RISCs → CMD BUFs (single aggregate arrow from middle of RISC stack)
connect(CX + CW, CY + 1.4, CB_X, CB_Y + CB_H/2, width=0.9, arrow=True)
# L1 ↔ NIU body (bidirectional data)
connect(LX + LW, LY + 2.0, NX, LY + 2.0, width=0.9, arrow=True)
connect(NX, LY + 2.3, LX + LW, LY + 2.3, width=0.9, arrow=True)
# CMD BUF → routers (via NIU-issued flits)
connect(CB_X + CB_W/2, NY + NH, 3.6, RTR_Y, width=0.9, arrow=True)
connect(CB_X + CB_W/2 + 0.2, NY + NH, 9.75, RTR_Y, width=0.9, arrow=True)
# Routers → mesh (down arrows out of tile)
for xr in (2.0, 5.2, 8.4, 11.4):
    connect(xr, RTR_Y + 0.6, xr, mesh_y, width=0.9, arrow=True)


# ==========================================================
# SAVE
# ==========================================================
out = (r'C:\Users\jbpar\AppData\Local\Temp\claude\C--NPU-Challenge'
       r'\44f2a027-4a0b-4b75-aa87-9f013a77983e\scratchpad'
       r'\blackhole_noc_figure.pptx')
prs.save(out)
print('OK ->', out)
